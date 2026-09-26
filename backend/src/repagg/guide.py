"""Import the r/RepTime "Who Makes the Best? Guide" spreadsheet as a scoring baseline.

Each row names a model (and often reference numbers), a movement, and up to three
factories ranked Best / 2nd / 3rd. Cell colour carries quality: yellow = NWBIG
(best of the best), green = Super Rep (good, not entirely perfect). Every listed
factory becomes a version with a *prior* (starting score) that the scoring engine
shrinks toward; real findings move it from there.
"""

import json
import re
import sqlite3
from datetime import date
from pathlib import Path

GUIDE_ID = "wmtb"
GUIDE_NAME = "r/RepTime Who Makes the Best? Guide"
GUIDE_URL = "https://www.reddit.com/r/RepTime/wiki/index/"

# Section order in the sheet (brand titles are logo images, so they're mapped by position).
SECTION_BRANDS = ["Rolex", "Patek Philippe", "Audemars Piguet", "Vacheron Constantin", "Richard Mille", "Omega",
                  "Cartier", "Tudor", "Jaeger-LeCoultre", "Panerai", "Breitling", "IWC"]
YELLOW = {"FFFFE599"}
GREEN = {"FFB6D7A8", "FFB7D7A8", "FF92D050"}

# Starting score (0..10) by quality flag and rank; weight in claim-weight units (~4 solid findings).
PRIORS = {("nwbig", 1): 9.0, ("nwbig", 2): 8.9, ("nwbig", 3): 8.8,
          ("super", 1): 8.2, ("super", 2): 8.1, ("super", 3): 8.0,
          (None, 1): 7.4, (None, 2): 7.1, (None, 3): 6.8}
PRIOR_WEIGHT = 4.0

NOT_A_FACTORY = re.compile(r"^(no longer available|unknown|unknown, just ask|custom|n/?a|-|\?)$", re.I)
VARIANT_WORDS = [("free sprung", "Free Sprung"), ("non weighted", "Non-weighted"), ("weighted", "Weighted"),
                 ("tungsten", "Tungsten"), ("heavy", "Heavy"), ("youth", "Youth")]

SCHEMA = """
CREATE TABLE IF NOT EXISTS guide_entry (
    id           INTEGER PRIMARY KEY,
    guide        TEXT NOT NULL,
    sheet_row    INTEGER NOT NULL,
    brand        TEXT NOT NULL,
    family       TEXT,
    model_text   TEXT,
    reference_id TEXT,
    movement     TEXT,
    rank         INTEGER NOT NULL,
    factory_raw  TEXT NOT NULL,
    factory_id   TEXT,
    build_id     TEXT,
    quality      TEXT,
    note         TEXT
);
CREATE TABLE IF NOT EXISTS baseline (
    build_id     TEXT PRIMARY KEY REFERENCES build(id),
    guide        TEXT NOT NULL,
    prior_mean   REAL NOT NULL,
    prior_weight REAL NOT NULL,
    rank         INTEGER NOT NULL,
    quality      TEXT,
    updated      TEXT NOT NULL
);
"""


def init_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(SCHEMA)


def parse_sheet(path: Path) -> tuple[str | None, list[dict]]:
    """Rows of the main sheet -> entries (one per listed factory). Returns (guide date, entries)."""
    import openpyxl

    ws = openpyxl.load_workbook(path).worksheets[0]
    updated = None
    for r in range(1, 30):
        for c in range(1, 8):
            v = ws.cell(r, c).value
            if hasattr(v, "date"):
                updated = v.date().isoformat()
    merged_a = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_col == 1:
            for r in range(rng.min_row, rng.max_row + 1):
                merged_a[r] = ws.cell(rng.min_row, 1).value
    headers = [r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").strip() == "Model Family"]
    entries = []
    for i, h in enumerate(headers[: len(SECTION_BRANDS)]):
        brand = SECTION_BRANDS[i]
        end = headers[i + 1] if i + 1 < len(headers) else ws.max_row + 1
        ctx_model, ctx_family = None, None
        for r in range(h + 1, end):
            family = _txt(ws.cell(r, 1).value or merged_a.get(r))
            b_raw = ws.cell(r, 2).value
            model_text = _txt(b_raw)
            movement = _txt(ws.cell(r, 3).value)
            if family and ("ORDER THESE" in family.upper() or family.lower().startswith(("want to", "note"))):
                continue
            if family != ctx_family:
                ctx_family, ctx_model = family, None
            cells = [ws.cell(r, c) for c in (4, 5, 6)]
            if not any(_txt(c.value) for c in cells):
                if model_text and not (model_text or "").lower().startswith("note"):
                    ctx_model = model_text  # header row for the sub-rows that follow (e.g. "5711 All Dial variants")
                continue
            indented = isinstance(b_raw, str) and b_raw.startswith(" ")
            if not model_text:
                model_text = ctx_model
            elif indented and ctx_model:
                model_text = f"{ctx_model} — {model_text}"
            else:
                ctx_model = model_text
            if "diw" in (model_text or "").lower():
                continue  # DIY builds, not factory versions
            for rank, cell in enumerate(cells, start=1):
                raw = _txt(cell.value)
                if not raw:
                    continue
                quality = _quality(cell)
                for name, variant, note in _factories(raw):
                    entries.append({
                        "sheet_row": r, "brand": brand, "family": family, "model_text": model_text, "movement": movement,
                        "rank": rank, "factory_raw": raw, "factory": name, "variant": variant, "quality": quality, "note": note,
                    })
    return updated, entries


def import_guide(conn: sqlite3.Connection, path: Path) -> dict:
    """Replace this guide's entries and baseline priors; create versions it names. Idempotent."""
    from . import scoring
    from .extract import _resolve_reference, _rescore
    from .versions import add_release, ensure_build, split_label

    init_schema(conn)
    updated, entries = parse_sheet(path)
    updated = updated or date.today().isoformat()
    conn.execute("DELETE FROM guide_entry WHERE guide = ?", (GUIDE_ID,))
    conn.execute("DELETE FROM baseline WHERE guide = ?", (GUIDE_ID,))
    best: dict[str, tuple] = {}
    skipped = 0
    for e in entries:
        refs = _references(e["brand"], e["model_text"] or "")
        family = _family(e)
        ref_ids = [rid for rid in (_resolve_reference(conn, e["brand"], ref, family, trusted=True) for ref in refs) if rid] \
            or [rid for rid in [_resolve_reference(conn, e["brand"], "", family, trusted=True)] if rid]
        factory_id = _factory(conn, e["factory"])
        if not ref_ids or not factory_id:
            skipped += 1
            continue
        for rid in ref_ids:
            release, variant = split_label(e["variant"])
            build_id = ensure_build(conn, factory_id, rid, variant, e["movement"])
            add_release(conn, build_id, release, updated, "named in the WMTB guide")
            conn.execute(
                "INSERT INTO guide_entry(guide, sheet_row, brand, family, model_text, reference_id, movement, rank, factory_raw, factory_id, build_id, quality, note) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (GUIDE_ID, e["sheet_row"], e["brand"], family, e["model_text"], rid, e["movement"], e["rank"], e["factory_raw"],
                 factory_id, build_id, e["quality"], e["note"]),
            )
            prior = PRIORS[(e["quality"], e["rank"])]
            if build_id not in best or prior > best[build_id][0]:
                best[build_id] = (prior, e["rank"], e["quality"])
    for build_id, (prior, rank, quality) in best.items():
        conn.execute("INSERT INTO baseline VALUES (?, ?, ?, ?, ?, ?, ?)", (build_id, GUIDE_ID, prior, PRIOR_WEIGHT, rank, quality, updated))
    # Tidy up what an earlier import created but this one no longer names (renamed factories, etc.).
    evidence = ("SELECT build_id FROM claim UNION SELECT build_id FROM qc_verdict UNION SELECT build_id FROM price_point "
                "UNION SELECT build_id FROM event WHERE build_id IS NOT NULL UNION SELECT build_id FROM baseline")
    stale = [r[0] for r in conn.execute(f"SELECT id FROM build WHERE id NOT IN ({evidence})")]
    for bid in stale:
        for table in ("release", "score", "tier", "defect"):
            conn.execute(f"DELETE FROM {table} WHERE build_id = ?", (bid,))
        conn.execute("UPDATE photo SET build_id = NULL WHERE build_id = ?", (bid,))
        conn.execute("DELETE FROM build WHERE id = ?", (bid,))
    for (fid,) in conn.execute("SELECT id FROM factory WHERE notes = 'from WMTB guide' AND id NOT IN (SELECT factory_id FROM build) "
                               "AND id NOT IN (SELECT factory_id FROM event WHERE factory_id IS NOT NULL)").fetchall():
        conn.execute("DELETE FROM factory_alias WHERE factory_id = ?", (fid,))
        conn.execute("DELETE FROM factory WHERE id = ?", (fid,))
    conn.execute("INSERT INTO meta(key, value) VALUES ('guide_wmtb', ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                 (json.dumps({"name": GUIDE_NAME, "url": GUIDE_URL, "updated": updated, "file": path.name}),))
    conn.commit()
    _rescore(conn)
    by_quality = {q: sum(1 for v in best.values() if v[2] == q) for q in ("nwbig", "super", None)}
    return {"guide_date": updated, "entries": len(entries), "versions": len(best), "skipped": skipped,
            "nwbig": by_quality["nwbig"], "super": by_quality["super"], "listed": by_quality[None], "tiers": _tier_counts(conn, scoring)}


def _tier_counts(conn, scoring) -> dict:
    latest = conn.execute("SELECT MAX(as_of) FROM tier").fetchone()[0]
    return dict(conn.execute("SELECT tier, COUNT(*) FROM tier WHERE as_of = ? GROUP BY tier", (latest,)).fetchall()) if latest else {}


def _quality(cell) -> str | None:
    f = cell.fill
    rgb = f.fgColor.rgb if f is not None and f.fill_type and f.fgColor is not None and f.fgColor.type == "rgb" else None
    return "nwbig" if rgb in YELLOW else "super" if rgb in GREEN else None


def _factories(raw: str) -> list[tuple[str, str | None, str | None]]:
    """'QF Free Sprung' -> ('QF', 'Free Sprung'); 'V7F/VSF' -> two; 'KVF - thick but best decorated' -> note."""
    text, note = raw, None
    if " - " in text:
        text, note = text.split(" - ", 1)
    notes = re.findall(r"\(([^)]*)\)", text)
    text = re.sub(r"\([^)]*\)", "", text)
    if notes:
        note = "; ".join([n for n in [note, *notes] if n])
    out = []
    for part in re.split(r"\s*/\s*", text):
        part = part.strip().rstrip("*").strip()
        if not part or NOT_A_FACTORY.match(part) or part.lower().startswith(("diw", "unknown", "no longer")):
            continue
        variant_bits = [label for word, label in VARIANT_WORDS if word in part.lower()]
        if "Non-weighted" in variant_bits and "Weighted" in variant_bits:
            variant_bits.remove("Weighted")
        m = re.search(r"\bv(\d+(?:\.\d+)?)\b", part, re.I)
        if m:
            variant_bits.insert(0, f"V{m.group(1)}")
        name = part
        for word, _ in VARIANT_WORDS:
            name = re.sub(word, " ", name, flags=re.I)
        name = re.sub(r"\bv\d+(?:\.\d+)?\b|\bfactory\b", " ", name, flags=re.I)
        name = re.sub(r"\s+", " ", name).strip(" -,")
        name = re.sub(r"(?<=\w) F$", "F", name)  # the sheet writes some as "3S F", "TC F"
        if name:
            out.append((name, " ".join(variant_bits) or None, note))
    return out


def _references(brand: str, text: str) -> list[str]:
    """Reference numbers named in a model cell."""
    if brand == "Richard Mille":
        m = re.search(r"\bRM\s?0?(\d{2,3})(?:-(\d{2}))?", text, re.I)
        return [f"RM {int(m.group(1)):03d}" + (f"-{m.group(2)}" if m.group(2) else "")] if m else []
    text = re.sub(r"(\d{5,6}) (BLRO|BLNR|VTNR|GRNR|LN|LV|LB)\b", r"\1\2", text)
    refs = re.findall(r"\b(\d{4,6}(?:/\d[A-Z]?)?[A-Z]{0,5})\b", text)
    min_digits = 4 if brand == "Patek Philippe" else 5  # Patek uses 4-digit references
    return list(dict.fromkeys(r for r in refs if len(re.sub(r"\D", "", r.split("/")[0])) >= min_digits))


def _family(e: dict) -> str:
    fam = (e["family"] or "").strip()
    fam = re.sub(r"\s*\((vintage|[^)]*)\)\s*$", "", fam)
    if not fam and e["brand"] == "Richard Mille":
        refs = _references("Richard Mille", e["model_text"] or "")
        fam = refs[0].split("-")[0] if refs else "Richard Mille"
    return {"GMT Master II": "GMT-Master II", "GMT Master": "GMT-Master", "Explorer I": "Explorer", "Air-king": "Air-King",
            "Yachtmaster": "Yacht-Master", "DayDate": "Day-Date", "Skydweller": "Sky-Dweller", "Seadweller": "Sea-Dweller",
            "Datejust 36": "Datejust", "Datejust 41": "Datejust", "Datejust 28": "Lady-Datejust", "Datejust 31": "Lady-Datejust",
            "Deville": "De Ville", "Panthere": "Panthère", "Rendezvous": "Rendez-Vous", "Master Ultra Thin": "Master",
            "Big Pilot's / Pilot's": "Pilot's Watch", "SuperOcean": "Superocean", "Seamaster": "Seamaster Diver 300M",
            "Planet Ocean": "Seamaster Planet Ocean"}.get(fam, fam)


def _factory(conn: sqlite3.Connection, name: str) -> str | None:
    """Known factory by alias (also trying the name with/without a trailing 'F'); otherwise add it as a
    confirmed factory from the guide."""
    if conn.execute("SELECT 1 FROM factory_block WHERE alias = ?", (name,)).fetchone():
        return None
    for key in (name, name + "F" if not name.upper().endswith("F") else name[:-1], name.replace(" ", "")):
        row = conn.execute("SELECT factory_id FROM factory_alias WHERE alias = ?", (key,)).fetchone()
        if row:
            return row[0]
    fid = re.sub(r"[^a-z0-9]+", "", name.lower().replace("+", "plus")) or None
    if not fid:
        return None
    conn.execute("INSERT OR IGNORE INTO factory(id, name, status, notes, needs_review) VALUES (?, ?, 'unknown', 'from WMTB guide', 0)", (fid, name))
    conn.execute("UPDATE factory SET name = ? WHERE id = ? AND notes = 'from WMTB guide'", (name, fid))  # pick up name clean-ups
    conn.execute("INSERT OR IGNORE INTO factory_alias VALUES (?, ?)", (name, fid))
    return fid


def _txt(v) -> str | None:
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None
